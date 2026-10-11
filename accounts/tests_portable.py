from io import BytesIO,StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from django.test import TestCase,Client,override_settings
from django.core.files.base import ContentFile
from django.core.management import call_command
from django.core.mail import get_connection
from django.core.exceptions import ImproperlyConfigured
from django.conf import settings
from openpyxl import load_workbook
from .models import User,Department
from documents.models import Document
class PortableTests(TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.override=override_settings(MEDIA_ROOT=Path(self.temp.name));self.override.enable();self.addCleanup(self.override.disable)
    def test_first_director_and_repeat(self):
        with patch('builtins.input',side_effect=['owner','owner@example.org','Test Dept']),patch('getpass.getpass',return_value='A-safe-test-passphrase-9927!'):
            call_command('initialize_instance',stdout=StringIO())
        user=User.objects.get();self.assertTrue(user.check_password('A-safe-test-passphrase-9927!'));self.assertTrue(user.is_director());self.assertTrue(user.is_superuser)
        with patch('builtins.input',side_effect=AssertionError('Should not prompt on populated instance')):call_command('initialize_instance',stdout=StringIO())
        self.assertEqual(User.objects.count(),1)
    def test_views_branding_and_empty_import(self):
        dept=Department.objects.create(name='Test');user=User.objects.create_user(username='director',role='director',department=dept,is_staff=True,is_superuser=True)
        client=Client();client.force_login(user)
        for path in ['/dashboard/','/tasks/','/tasks/kanban/','/events/meropriyatiya/','/events/meropriyatiya/import/','/documents/','/wiki/','/analytics/','/schedule/','/admin/']:
            r=client.get(path);self.assertIn(r.status_code,[200,302],path)
        r=client.get('/events/meropriyatiya/import/template/');self.assertEqual(r.status_code,200)
        wb=load_workbook(BytesIO(r.content));self.assertEqual(wb['Мероприятия'].max_row,1)
        self.assertContains(client.get('/dashboard/'),settings.SITE_NAME)
    def test_files_private_and_permissions(self):
        a=Department.objects.create(name='A');b=Department.objects.create(name='B')
        owner=User.objects.create_user(username='owner',role='head',department=a)
        other=User.objects.create_user(username='other',role='head',department=b)
        doc=Document.objects.create(title='Secret',department=a,category='order',uploaded_by=owner)
        doc.file.save('Пример.txt',ContentFile(b'private content'))
        url=doc.file.url;self.assertIn('/private-files/',url)
        client=Client();self.assertEqual(client.get(url).status_code,302)
        client.force_login(other);self.assertEqual(client.get(url).status_code,404)
        client.force_login(owner);r=client.get(url);self.assertEqual(r.status_code,200);self.assertEqual(b''.join(r.streaming_content),b'private content')
        self.assertEqual(client.get(url.split('?')[0]).status_code,404)
        self.assertEqual(client.get('/media/'+doc.file.name).status_code,404)
    @override_settings(EMAIL_BACKEND="accounts.mail_backend.DisabledEmailBackend")
    def test_smtp_disabled(self):
        with self.assertRaises(ImproperlyConfigured):get_connection().send_messages([])
