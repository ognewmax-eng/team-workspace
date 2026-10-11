from io import BytesIO
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import render
from django.views.decorators.http import require_safe
from django.utils import timezone
from .reporting import AnalyticsFilterForm, report


def context(request):
    params = request.GET.copy()
    if not params.get('period'): params['period'] = '30'
    form = AnalyticsFilterForm(params)
    valid = form.is_valid()
    if valid: data = form.cleaned_data
    else:
        default = AnalyticsFilterForm({'period': '30'}); default.is_valid(); data = default.cleaned_data
    result = report(data)
    result.update(form=form, export_query=params.urlencode())
    return result, valid


@login_required
@require_safe
def analytics_dashboard_view(request):
    if not request.user.is_director(): return HttpResponseForbidden('Аналитика доступна директору.')
    ctx, valid = context(request)
    return render(request, 'analytics/dashboard.html', ctx, status=200 if valid else 400)


@login_required
@require_safe
def analytics_export_view(request):
    if not request.user.is_director(): return HttpResponseForbidden('Аналитика доступна директору.')
    ctx, valid = context(request)
    if not valid: return render(request, 'analytics/dashboard.html', ctx, status=400)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    def sheet(title, headers, rows):
        ws = wb.create_sheet(title); ws.append(headers)
        for row in rows:
            ws.append(row)
            # User-entered names and titles remain literal strings, including = formulas.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str): cell.data_type = 's'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='6254C7')
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for column in range(1, len(headers) + 1): ws.column_dimensions[get_column_letter(column)].width = 26
        return ws
    s = ctx['stats']
    sheet('Показатели', ['Показатель', 'Значение'], [
        ['Отдел', ctx['department_name']], ['Начало периода', ctx['start'].isoformat()], ['Конец периода', ctx['end'].isoformat()],
        ['Неархивные задачи, кроме отменённых мероприятий', s['total']], ['Незавершённые сейчас', s['active']],
        ['Просроченные сейчас', s['overdue']], ['На проверке сейчас', s['pending']],
        ['Созданы за период', s['created_period']], ['Завершены за период', s['completed_period']],
        ['Без срока сейчас', s['no_deadline']], ['Без исполнителя сейчас', s['unassigned']],
        ['Мероприятия с датой проведения в периоде', ctx['event_stats']['period']],
        ['Из них отмечены завершёнными', ctx['event_stats']['completed']],
        ['Из них отменены', ctx['event_stats']['cancelled']],
        ['Предстоящие неотменённые мероприятия', ctx['event_stats']['upcoming']],
        ['Принцип', 'Состояния — на момент выгрузки; период влияет на создание, завершение и назначения.'],
        ['Совместные задачи', 'Учитываются в каждом своём отделе. Строки отделов не суммируются.'],
        ['Назначения', 'Личный завершённый статус исполнителя; это не количество целых задач.']])
    sheet('Отделы', ['Отдел', 'Сотрудники', 'Задачи', 'Незавершённые', 'Завершённые', 'Просроченные', 'Выполнение %'],
        [[r['department'].name, r['members'], r['total'], r['active'], r['completed'], r['overdue'], r['completion_rate']] for r in ctx['department_stats']])
    sheet('Назначения', ['Сотрудник', 'Отдел', 'Активные назначения', 'Просроченные', 'Завершены за период'],
        [[p.get_full_name() or p.username, p.department.name if p.department else '', p.active_count, p.overdue_count, p.completed_count] for p in ctx['people']])
    sheet('Динамика', ['Дата', 'Созданы задачи', 'Завершены задачи'], [[d['day'].isoformat(), d['created'], d['completed']] for d in ctx['daily']])
    sheet('Просроченные', ['Задача', 'Основной отдел', 'Срок', 'Статус'],
        [[t.title, t.department.name, timezone.localtime(t.deadline).isoformat(), t.get_status_display()] for t in ctx['blockers']])
    buffer = BytesIO(); wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="yamolod_analytics_%s_%s.xlsx"' % (ctx['start'], ctx['end'])
    response['Cache-Control'] = 'private, no-store'
    return response
