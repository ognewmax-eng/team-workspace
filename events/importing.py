"""No JS execution: accept bounded Excel rows or the colleague's JSON array."""
from datetime import datetime,timedelta
from decimal import Decimal,InvalidOperation
from hashlib import sha256
from pathlib import Path
import json,re,zipfile
from django.core.exceptions import ValidationError
from django.db import transaction,IntegrityError
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from accounts.models import Department,User
from .models import Meropriyatie,EventPlan,EventWorkspace,EventBudget,EventImportBatch,EventImportRecord
from .collaboration import log
from .planning import period

HEADERS=[('source_code','Код источника'),('title','Название'),('year','Год'),('month_start','Месяц начала'),
    ('month_end','Месяц окончания'),('event_date','Дата и время'),('event_type','Тип'),('organizer','Организатор'),
    ('department','Отдел'),('responsible','Логин ответственного'),('location','Место'),('description','Описание'),
    ('participants_target','План участников'),('allocated','Лимит бюджета')]
MAX_ROWS=1000


def can_import(user):return user.is_director() or (user.is_head() and bool(user.department_id))


def departments(user):
    return Department.objects.all() if user.is_director() else Department.objects.filter(pk=user.department_id)


def norm(value):return ' '.join(str(value or '').split()).casefold()


def text(value,limit,label):
    if value is None:return ''
    result=str(value).strip()
    if len(result)>limit:raise ValidationError('%s: не более %s символов.'%(label,limit))
    if '\x00' in result:raise ValidationError('%s: недопустимый символ.'%label)
    return result


def integer(value,low,high,label,optional=False):
    if optional and value in ['',None]:return None
    try:
        number=Decimal(str(value))
        if not number.is_finite() or number!=int(number) or not low<=number<=high:raise ValueError()
        return int(number)
    except (ValueError,TypeError,InvalidOperation,OverflowError):raise ValidationError('%s: целое число от %s до %s.'%(label,low,high)) from None


def date_value(value):
    if value in ['',None]:return None
    if isinstance(value,datetime):result=value
    else:
        raw=str(value).strip()
        try:result=parse_datetime(raw)
        except ValueError:result=None
        if result is None:
            for fmt in ['%d.%m.%Y %H:%M','%Y-%m-%d %H:%M','%d.%m.%Y']:
                try:result=datetime.strptime(raw,fmt);break
                except ValueError:pass
        if result is None:raise ValidationError('Дата: используйте ДД.ММ.ГГГГ ЧЧ:ММ или дату Excel.')
    if not 2000<=result.year<=2100:raise ValidationError('Дата: год от 2000 до 2100.')
    return timezone.make_aware(result) if timezone.is_naive(result) else result


def money(value):
    if value in ['',None]:return None
    try:
        result=Decimal(str(value).replace(' ','').replace('\xa0','').replace(',','.'))
        if not result.is_finite() or result<0 or result>Decimal('999999999999.99') or result!=result.quantize(Decimal('0.01')):raise ValueError()
        return result
    except (InvalidOperation,ValueError):raise ValidationError('Лимит бюджета: неотрицательная сумма, не более двух знаков после запятой.') from None


def read_upload(upload):
    if upload.size>5*1024*1024:raise ValidationError('Файл должен быть не больше 5 МБ.')
    suffix=Path(upload.name).suffix.lower()
    if suffix=='.xlsx':
        try:
            with zipfile.ZipFile(upload) as archive:
                if len(archive.infolist())>2000 or sum(x.file_size for x in archive.infolist())>40*1024*1024:
                    raise ValidationError('Распакованный Excel слишком велик.')
            upload.seek(0)
            from openpyxl import load_workbook
            wb=load_workbook(upload,read_only=True,data_only=False,keep_links=False)
            try:
                if 'Мероприятия' not in wb.sheetnames:raise ValidationError('Не найден лист «Мероприятия». Скачайте шаблон.')
                ws=wb['Мероприятия'];result=[]
                if ws.max_column and ws.max_column>len(HEADERS):raise ValidationError('В листе есть лишние столбцы. Используйте структуру шаблона.')
                rows=ws.iter_rows(max_col=len(HEADERS)+1)
                header=next(rows)
                if [c.value for c in header[:len(HEADERS)]]!=[label for key,label in HEADERS] or header[-1].value is not None:raise ValidationError('Заголовки не совпадают с шаблоном.')
                for number,row in enumerate(rows,2):
                    if number>MAX_ROWS+1:raise ValidationError('Не более %s мероприятий за одну загрузку.'%MAX_ROWS)
                    if not any(c.value is not None for c in row):continue
                    if row[-1].value is not None:raise ValidationError('Строка %s: лишний столбец.'%number)
                    if any(c.data_type=='f' for c in row):raise ValidationError('Строка %s: формулы не поддерживаются. Вставьте значения.'%number)
                    record={key:(c.value.isoformat() if isinstance(c.value,datetime) else c.value) for (key,label),c in zip(HEADERS,row)}
                    record['_row']=number;result.append(record)
                if ws.max_row and ws.max_row>MAX_ROWS+1:
                    # The bounded iterator must not silently drop later records.
                    raise ValidationError('Лист содержит больше %s строк. Удалите лишние строки или разделите файл.'%MAX_ROWS)
                return result
            finally:wb.close()
        except ValidationError:raise
        except Exception:raise ValidationError('Не удалось прочитать Excel. Проверьте, что это обычный файл .xlsx по шаблону.') from None
    if suffix in ['.html','.htm']:
        try:
            source=upload.read().decode('utf-8-sig')
            match=re.search(r'const\s+INITIAL_DATA\s*=\s*(\[.*?\]);',source,re.S)
            if not match:raise ValueError()
            data=json.loads(match.group(1));result=[]
            from .planning import MONTHS
            names=[m.lower() for m in MONTHS]
            if not isinstance(data,list) or len(data)>MAX_ROWS:raise ValueError()
            for record in data:
                if not isinstance(record,list) or len(record)!=5:raise ValueError()
                code,title,months,quarter,organizer=record
                parts=str(months).replace('–','-').replace('—','-').lower().split('-')
                start=names.index(parts[0].strip())+1;end=names.index(parts[-1].strip())+1
                result.append(dict(source_code='colleague-2026-'+str(code),title=title,year=2026,month_start=start,month_end=end,
                    organizer=organizer,description='Из HTML коллеги. Исходный период: %s. Исходный квартал: %s.'%(months,quarter),event_type='mz'))
            if not re.search(r'2026',source[:10000]):raise ValueError()
            return result
        except (ValueError,TypeError,UnicodeError):raise ValidationError('Поддерживается HTML дашборда 2026 с массивом INITIAL_DATA. JavaScript не выполняется.') from None
    raise ValidationError('Используйте .xlsx по шаблону или HTML дашборда коллеги.')


def prepare(raw,user,default_department):
    if not can_import(user) or not departments(user).filter(pk=default_department.pk).exists():raise ValidationError('Недоступен отдел импорта.')
    title=text(raw.get('title'),300,'Название')
    if not title:raise ValidationError('Укажите название.')
    date=date_value(raw.get('event_date'))
    if date:
        local=timezone.localtime(date);year,start,end=local.year,local.month,local.month
    else:
        year=integer(raw.get('year'),2000,2100,'Год');start=integer(raw.get('month_start'),1,12,'Месяц начала')
        end=integer(raw.get('month_end') if raw.get('month_end') not in ['',None] else start,1,12,'Месяц окончания')
        if end<start:raise ValidationError('Месяц окончания не может быть раньше начала.')
    department=default_department
    department_name=text(raw.get('department'),300,'Отдел')
    if department_name:
        matches=[d for d in departments(user) if norm(d.name)==norm(department_name)]
        if len(matches)!=1:raise ValidationError('Отдел не найден, неоднозначен или недоступен: '+department_name)
        department=matches[0]
    login=text(raw.get('responsible'),150,'Логин ответственного');responsible=None
    if login:
        responsible=User.objects.filter(username=login,is_active=True).first()
        if not responsible:raise ValidationError('Активный сотрудник с таким логином не найден.')
    kind=str(raw.get('event_type') or 'mz').strip()
    kind={'МЗ':'mz','Внеплановое':'vneplan'}.get(kind,kind)
    if kind not in dict(Meropriyatie.EVENT_TYPE_CHOICES):raise ValidationError('Тип: mz (МЗ) или vneplan (внеплановое).')
    data=dict(title=title,event_date=date.isoformat() if date else '',year=year,month_start=start,month_end=end,
        department_id=department.pk,responsible_id=responsible.pk if responsible else None,event_type=kind,
        organizer=text(raw.get('organizer'),300,'Организатор'),location=text(raw.get('location'),200,'Место'),
        description=text(raw.get('description'),20000,'Описание'),participants_target=integer(raw.get('participants_target'),0,10000000,'План участников',True),
        allocated=str(money(raw.get('allocated'))) if raw.get('allocated') not in ['',None] else None)
    code=text(raw.get('source_code'),120,'Код источника')
    natural=(department.pk,norm(title),year,start,end)
    key=sha256(json.dumps([department.pk,code] if code else natural,ensure_ascii=False).encode()).hexdigest()
    return data,key,natural


def preview(batch,user):
    prepared=[];seen=set();naturals=set()
    existing={(e.department_id,norm(e.title),*period(e)) for e in Meropriyatie.objects.filter(department__in=departments(user)).select_related('plan')}
    for index,raw in enumerate(batch.rows):
        entry={'index':index,'number':raw.get('_row',index+2),'title':str(raw.get('title') or '')[:300],'errors':[],'warnings':[],'duplicate':False}
        try:
            data,key,natural=prepare(raw,user,batch.department)
            entry.update(data=data,key=key,natural=natural)
            entry['duplicate']=key in seen or natural in naturals or natural in existing or EventImportRecord.objects.filter(key=key).exists()
            seen.add(key);naturals.add(natural)
            if not data['event_date']:entry['warnings'].append('Точная дата и время уточняются')
            if not data['responsible_id']:entry['warnings'].append('Ответственный сотрудник не назначен')
            if data['month_start']!=data['month_end'] and (data['month_start']-1)//3!=(data['month_end']-1)//3:entry['warnings'].append('Период охватывает несколько кварталов')
        except ValidationError as exc:entry['errors']=exc.messages
        prepared.append(entry)
    return prepared


def commit(batch,user,indices):
    with transaction.atomic():
        stored=EventImportBatch.objects.select_for_update().select_related('department').get(pk=batch.pk,actor=user)
        if stored.committed:return stored.result
        if stored.created_at<timezone.now()-timedelta(hours=24):raise ValidationError('Предпросмотр устарел. Загрузите файл заново.')
        if not indices:raise ValidationError('Выберите хотя бы одно мероприятие.')
        entries=preview(stored,user)
        if any(i<0 or i>=len(entries) for i in indices):raise ValidationError('Некорректный выбор строк.')
        selected=[entries[i] for i in sorted(indices)]
        if any(e['errors'] for e in selected):raise ValidationError('В выбранных строках есть ошибки. Исправьте файл или исключите строки.')
        created=[];skipped=0
        for entry in selected:
            if entry['duplicate']:skipped+=1;continue
            data=entry['data']
            event=Meropriyatie.objects.create(title=data['title'],description=data['description'],event_type=data['event_type'],
                event_date=date_value(data['event_date']),department_id=data['department_id'],responsible_id=data['responsible_id'],created_by=user)
            EventPlan.objects.create(event=event,year=data['year'],month_start=data['month_start'],month_end=data['month_end'],organizer=data['organizer'],
                participants_target=data['participants_target'],source_label=stored.source_label)
            EventWorkspace.objects.create(event=event,location=data['location'])
            if data['allocated'] is not None:EventBudget.objects.create(event=event,allocated=Decimal(data['allocated']))
            EventImportRecord.objects.create(event=event,batch=stored,key=entry['key'])
            log(event,user,'Добавлено из импорта: '+stored.source_label)
            created.append(event.pk)
        stored.committed=True;stored.result={'created':created,'skipped':skipped};stored.save(update_fields=['committed','result'])
    # Calendar sync happens only after the database transaction is committed.
    from .work_views import sync_event
    for event in Meropriyatie.objects.filter(pk__in=created,event_date__isnull=False):sync_event(event)
    return stored.result
