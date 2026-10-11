"""Application-generated exports. Explicit text cells prevent Excel formulas."""
from decimal import Decimal
from io import BytesIO
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from .planning import period,period_label,budget_totals


def sheet(workbook,name,headers,rows):
    ws=workbook.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='4C368A');cell.alignment=Alignment(wrap_text=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(55,max(15,max(len(str(c.value or '')) for c in col)+2))
        for cell in col[1:]:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(cell.value,(Decimal,float)):cell.number_format='#,##0.00'
    ws.row_dimensions[1].height=35
    return ws


def response(wb,name):
    output=BytesIO();wb.save(output)
    r=HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    r['Content-Disposition']='attachment; filename="%s"'%name
    return r


def export_budget(event):
    from .models import EventBudget
    budget=EventBudget.objects.filter(event=event).first();totals=budget_totals(budget)
    wb=Workbook();wb.remove(wb.active)
    sheet(wb,'Бюджет',['Показатель','Значение'],[
        ['Мероприятие',event.title],['Период',period_label(event)],['Источник финансирования',budget.funding_source if budget else ''],
        ['Согласование',budget.get_status_display() if budget else 'Нет сметы'],['Лимит, ₽',totals['allocated']],
        ['Смета, ₽',totals['estimate']],['Закупки / обязательства, ₽',totals['actual']],['Оплачено, ₽',totals['paid']],
        ['К оплате, ₽',totals['unpaid']],['Остаток лимита после закупок, ₽',totals['remaining']]])
    sheet(wb,'Позиции',['Наименование','Категория','Количество','Единица','Цена, ₽','Смета, ₽','Получено','Закупка, ₽','Оплачено, ₽','Статус','Поставщик','Счёт / договор','Поставка до','Примечание'],
        [[l.title,l.category,l.quantity,l.unit,l.unit_price,l.estimate,l.purchased_quantity,l.actual_amount,l.paid_amount,l.get_status_display(),l.supplier,l.document_number,l.due_date,l.note] for l in budget.lines.all()] if budget else [])
    sheet(wb,'Согласования',['Дата','Согласовал','Лимит, ₽','Смета, ₽','Закупки, ₽','Оплачено, ₽'],
        [[timezone.localtime(a.created_at).strftime('%Y-%m-%d %H:%M'),a.approver.get_full_name() or a.approver.username if a.approver else 'Удалённый пользователь',Decimal(a.snapshot['allocated']) if a.snapshot.get('allocated') is not None else None,Decimal(a.snapshot.get('totals',{}).get('estimate','0')),Decimal(a.snapshot.get('totals',{}).get('actual','0')),Decimal(a.snapshot.get('totals',{}).get('paid','0'))] for a in budget.approvals.select_related('approver').order_by('pk')] if budget else [])
    return response(wb,'event_%s_budget.xlsx'%event.pk)


def export_events(events):
    from .importing import HEADERS
    wb=Workbook();wb.remove(wb.active)
    rows=[]
    for event in events:
        year,start,end=period(event);plan=getattr(event,'plan',None);workspace=getattr(event,'workspace',None)
        rows.append(['site-event-'+str(event.pk),event.title,year,start,end,timezone.localtime(event.event_date).replace(tzinfo=None) if event.event_date else None,
            event.event_type,plan.organizer if plan else '',event.department.name,event.responsible.username if event.responsible else '',
            workspace.location if workspace else '',event.description,plan.participants_target if plan else None,None])
    ws=sheet(wb,'Мероприятия',[label for key,label in HEADERS],rows)
    for cell in ws['F'][1:]:cell.number_format='dd.mm.yyyy hh:mm'
    # Financial limits are intentionally exported only through the gated finance endpoint.
    return response(wb,'events_plan.xlsx')
